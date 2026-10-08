from pathlib import Path
from collections import Counter, defaultdict
from copy import copy
import csv, hashlib, json, re, shutil, unicodedata
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/"work"/"ogj2014_provenance"; OUT.mkdir(parents=True,exist_ok=True)
CANONICAL=ROOT/"src/notebooks/ml_data/OGJ2014_extracted_vs_Table1.xlsx"
MIRROR=ROOT/"data/OGJ2014_extracted_vs_Table1.xlsx"
TARGET="OGJ2014_Extracted_Rows"
NEW_COLS=["record_id","source_report_group_id","publication_id","publication_title","publication_id_status","field_id","field_name_normalized","field_id_status","project_id","case_id","case_id_status","tagging_confidence","tagging_rationale","needs_review"]
COMPANIES=sorted(set([
"Zulia PDVSA E&P","Anzoategui PDVSA E&P","Campo Jobo PDVSA E&P","Campo Pilon PDVSA E&P","Campo Cerro Negro PDVSA E&P","Maturin, Campo Mulata PDVSA E&P","Maturin, Campo Furrial PDVSA E&P","Lower Saxony, Grafschaft Bentheim Wintershall Holding AG","Lower Saxony, Diepholz Wintershall Holding AG","Rio Grande do Norte (RN) Petrobras","Bentheim Wintershall Holding AG","Point Fortin Petrotrin","Forest Reserve Petrotrin","Oropouche Petrotrin","Guapo Petrotrin","Pekanbaru PT Caltex","Espirito Santo Petrobras","Brazil Petrobras","Bahia Petrobras","Apache Canada","ExxonMobil Oil Canada","Shell Canada","Penn West Energy Trust","Pengrowth Corporation","Derek Oil & Gas Corp.","Derek Oil","MegaWest Energy Corp.","Encore Acquisition","Continental Resources","Chaparral Energy","Core Energy","Denbury Resources","Devon Canada","Devon","Energen Resources","ExxonMobil","Fasken","Hess","Kinder Morgan","Merit Energy","Occidental","ConocoPhillips","Husky Oil","Imperial Oil","CNRL","Cenovus","Penn West Energy","Seneca Resources","Tidelands","Great Western Drilling","Stanberry Oil","XTO Energy Inc.","XTO Energy","BP","Aera Energy","Berry","Carrizo","Chevron","Anadarko","Breitburn Energy","Apache","Orla Petco","Naftex","Bayou State","Whiting Petroleum","Trinity","Stockdale","Petrobras","Petrotrin","Wintershall Holding AG","PT Caltex"]),key=len,reverse=True)
GEO=set("Calif. Calif Tex. Tex Wyo. Wyo Mo. Mo Mont. Mont Miss. Miss Alta. Alta Sask. Sask La. La Okla. Okla Alas. Alas Mich. Mich NM N.Mex. N.M. ND N.Dak. SD S.Dak. MT TX OCS Fla. Pa. Colo. Kan. Neb. Utah Ariz. Nev. Ark. Ala. Tenn. Ga. Ill. Ind. Ohio Va. N.Y. N.J. Alberta Saskatchewan".split())
DATE_RE=re.compile(r"(?<!\d)(?:\d{1,2}/\d{1,2}(?:/\d{1,2})?(?:-\d{1,2}/\d{1,2})?|19\d{2}|20\d{2})(?!\d)")
TECH=["Miscible acid gas","Miscible CO2","Miscible HC","CO2 immiscible","CO2 miscible","Hot water","Enhanced cyclic steam","Technology","Surfactants","Nitrates","Microbial","Polymer","Combustion","Steam"]

def ns(v):
    if v is None:return ""
    return re.sub(r"\s+"," ",unicodedata.normalize("NFKC",str(v)).replace("\u00a0"," ")).strip()
def nk(v):return ns(v).casefold()
def hid(p,k):return p+"-"+hashlib.sha256(k.encode()).hexdigest()[:12]
def source_norm(v):return ns(v)
def find_company(raw):
    low=raw.casefold()
    for c in COMPANIES:
        p=low.find(c.casefold())
        if p>=0 and not (c=="Resources" and p!=0):return c,p
    return None,None
def striptech(s):
    s=ns(s)
    s=re.sub(r"^TECHNOLOGY","",s,flags=re.I).strip()
    changed=True
    while changed:
        changed=False
        for t in sorted(TECH,key=len,reverse=True):
            if s.casefold().startswith(t.casefold()+" "):
                s=s[len(t):].strip();changed=True;break
    return s
def parse_field(raw):
    raw=ns(raw); company,pos=find_company(raw)
    if company:
        rest=raw[pos+len(company):].strip(); m=DATE_RE.search(rest)
        if not m:return None,"no_date",""
        pre=striptech(rest[:m.start()]); toks=pre.split(); gi=None
        for i,t in enumerate(toks):
            if t in GEO:gi=i;break
        if gi is not None: field=" ".join(toks[:gi]).strip(" ,"); geo=toks[gi]
        else: field=pre.strip(" ,"); geo=""
        if field and len(field)<=100:return field,"organization_field_span",geo
        return None,"ambiguous",""
    m=DATE_RE.search(raw)
    if m:
        pre=ns(raw[:m.start()]).strip(" ,")
        if 1<=len(pre.split())<=4 and not re.search(r"\b(?:Steam|Polymer|Microbial|Combustion|miscible|immiscible)\b",pre,re.I):
            return pre,"standalone_field", ""
    return None,"unresolved",""

def matrix(path):
    wb=load_workbook(path,data_only=False,read_only=False)
    return wb,{ws.title:[[c.value for c in row] for row in ws.iter_rows()] for ws in wb.worksheets}

def verify(before,after):
    fails=[]
    if list(before)!=list(after):fails.append("sheet names/order changed")
    for s in before:
        if s==TARGET:
            a,b=before[s],after[s]
            for r in range(len(a)):
                for c in range(15):
                    av=a[r][c] if c<len(a[r]) else None; bv=b[r][c] if c<len(b[r]) else None
                    if av!=bv:fails.append(f"{s}!R{r+1}C{c+1} changed");return fails
        elif before[s]!=after[s]:fails.append(f"{s} original values changed");return fails
    return fails

def main():
    wb0,ma=matrix(CANONICAL); wb1,mb=matrix(MIRROR)
    copy_diff=[] 
    if list(ma)!=list(mb):copy_diff.append("sheet names/order differ")
    else:
        for s in ma:
            if ma[s]!=mb[s]:copy_diff.append(f"sheet {s} differs")
    if copy_diff:raise SystemExit("SOURCE COPIES DIFFER: "+"; ".join(copy_diff))
    ws=wb0[TARGET]; headers=[c.value for c in ws[1]]
    expected=["technique","formation_category","porosity_min_pct","porosity_max_pct","perm_min_md","perm_max_md","depth_min_ft","depth_max_ft","api_min","api_max","visc_min_cp","visc_max_cp","so_start_min_pct","so_start_max_pct","raw_line"]
    if headers!=expected:raise SystemExit(f"Unexpected headers: {headers}")
    raw_i=14; sr_i=headers.index("source_report") if "source_report" in headers else None
    rows=[(r,[ws.cell(r,c).value for c in range(1,16)]) for r in range(2,ws.max_row+1)]
    occ=Counter(); raw_counts=Counter()
    tags=[]
    for er,vals in rows:
        content="\x1f".join("" if v is None else str(v) for v in vals);occ[content]+=1
        rid=hid("REC",content)+f"-{occ[content]:02d}"
        raw=vals[raw_i]; rn=source_norm(raw); raw_counts[rn]+=1
        srg=hid("SRG",source_norm(vals[sr_i])) if sr_i is not None and source_norm(vals[sr_i]) else ""
        field,why,geo=parse_field(raw or "")
        fid=hid("FIELD",nk(field)+"|"+nk(geo)) if field else ""
        fstat="provisional" if field else "needs_review"
        cid=hid("CASE",rn) if rn else ""
        cstat="provisional" if cid else "needs_review"
        conf="Low" if not field else ("High" if raw_counts[rn]>1 else "Medium")
        review="Yes" if fstat=="needs_review" or cstat=="needs_review" else "No"
        rat=[]
        rat.append("No source_report column exists; source_report_group_id was not derived from raw_line.")
        rat.append("No explicit publication citation/title is present; publication identity was not invented.")
        rat.append(f"Field candidate '{field}' extracted conservatively from raw_line." if field else "Field identity unresolved from workbook evidence.")
        rat.append("case_id is a deterministic hash of exact normalized raw_line; repeated identical raw_line values share the case group." if cid else "No reliable case evidence.")
        tags.append({"excel_row":er,"source_report":vals[sr_i] if sr_i is not None else "","record_id":rid,"source_report_group_id":srg,"publication_id":"","publication_title":"","publication_id_status":"needs_review","field_id":fid,"field_name_normalized":ns(field),"field_id_status":fstat,"project_id":"","case_id":cid,"case_id_status":cstat,"tagging_confidence":conf,"tagging_rationale":" ".join(rat),"needs_review":review})
    for src,name in [(CANONICAL,"OGJ2014_extracted_vs_Table1_provenance.xlsx"),(MIRROR,"OGJ2014_extracted_vs_Table1_provenance_mirror.xlsx")]:
        out=OUT/name;shutil.copy2(src,out);wb=load_workbook(out,data_only=False,read_only=False);ws=wb[TARGET];start=ws.max_column+1;lh=ws.cell(1,start-1)
        for j,h in enumerate(NEW_COLS):
            c=ws.cell(1,start+j,h);c._style=copy(lh._style);c.font=copy(lh.font);c.fill=copy(lh.fill);c.border=copy(lh.border);c.alignment=copy(lh.alignment)
        for i,t in enumerate(tags,2):
            for j,h in enumerate(NEW_COLS):ws.cell(i,start+j,t[h])
        wb.save(out)
    with (OUT/"ogj2014_row_tagging_audit.csv").open("w",newline="",encoding="utf-8-sig") as f:
        fields=["excel_row","source_report"]+NEW_COLS;w=csv.DictWriter(f,fieldnames=fields);w.writeheader();w.writerows(tags)
    mp={}
    for t in tags:
        for typ,ident,key,evid,stat in [("record_id",t["record_id"],"row-content-hash","Original row values hashed with deterministic occurrence counter.","confirmed"),("case_id",t["case_id"],"exact-normalized-raw_line","Exact normalized raw_line text.","provisional"),("field_id",t["field_id"],nk(t["field_name_normalized"]),"Named field candidate extracted from raw_line.","provisional")]:
            if not ident:continue
            if ident not in mp:mp[ident]={"id_type":typ,"id":ident,"normalized_key":key,"supporting_evidence":evid,"status":stat,"confidence":t["tagging_confidence"],"source_rows":[],"needs_review":"No"}
            mp[ident]["source_rows"].append(str(t["excel_row"]))
            if t["needs_review"]=="Yes":mp[ident]["needs_review"]="Yes"
    with (OUT/"ogj2014_id_mapping.csv").open("w",newline="",encoding="utf-8-sig") as f:
        fs=["id_type","id","normalized_key","supporting_evidence","status","confidence","source_rows","needs_review"];w=csv.DictWriter(f,fieldnames=fs);w.writeheader()
        for x in sorted(mp.values(),key=lambda z:(z["id_type"],z["id"])):x["source_rows"]=";".join(x["source_rows"]);w.writerow(x)
    ver={}
    for src,name in [(CANONICAL,"OGJ2014_extracted_vs_Table1_provenance.xlsx"),(MIRROR,"OGJ2014_extracted_vs_Table1_provenance_mirror.xlsx")]:
        _,before=matrix(src);_,after=matrix(OUT/name);outwb=load_workbook(OUT/name,data_only=False,read_only=False);outws=outwb[TARGET]
        ver[name]={"failures":verify(before,after),"original_rows":len(before[TARGET]),"output_rows":len(after[TARGET]),"original_columns":15,"output_columns":29,"headers_at_end":[outws.cell(1,c).value for c in range(16,30)]==NEW_COLS,"technique_unchanged":all(before[TARGET][r][0]==after[TARGET][r][0] for r in range(len(before[TARGET])))}
    summary={"source_copies_identical":True,"rows_tagged":len(tags),"source_report_column_present":sr_i is not None,"distinct_source_report_groups":len({t["source_report_group_id"] for t in tags if t["source_report_group_id"]}),"distinct_record_ids":len({t["record_id"] for t in tags}),"distinct_case_ids":len({t["case_id"] for t in tags if t["case_id"]}),"distinct_field_ids":len({t["field_id"] for t in tags if t["field_id"]}),"distinct_project_ids":0,"publication_confirmed":0,"publication_provisional":0,"publication_unknown":0,"publication_needs_review":len(tags),"field_confirmed":0,"field_provisional":sum(t["field_id_status"]=="provisional" for t in tags),"field_unknown":0,"field_needs_review":sum(t["field_id_status"]=="needs_review" for t in tags),"case_confirmed":0,"case_provisional":sum(t["case_id_status"]=="provisional" for t in tags),"case_unknown":0,"case_needs_review":sum(t["case_id_status"]=="needs_review" for t in tags),"rows_needing_review":sum(t["needs_review"]=="Yes" for t in tags),"ambiguous_or_conflicting_assignments":sum(t["field_id_status"]=="needs_review" for t in tags),"exact_raw_line_groups":len([x for x in raw_counts if x]),"repeated_exact_raw_line_groups":sum(v>1 for v in raw_counts.values()),"verification":ver}
    (OUT/"ogj2014_provenance_summary.json").write_text(json.dumps(summary,indent=2),encoding="utf-8")
    (OUT/"ogj2014_provenance_summary.md").write_text("# OGJ2014 provenance preparation\n\n"+json.dumps(summary,indent=2)+"\n\nAll added IDs are metadata only and must not be used as model features. The workbook contains no source_report column or explicit publication citation, so those identities were left unresolved rather than invented. raw_line was not repurposed as source_report.\n",encoding="utf-8")
    print(json.dumps(summary,indent=2))
if __name__=="__main__":main()
