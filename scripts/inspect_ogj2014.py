from __future__ import annotations
import json, re
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
paths = [
    ROOT / "data/OGJ2014_extracted_vs_Table1.xlsx",
    ROOT / "src/notebooks/ml_data/OGJ2014_extracted_vs_Table1.xlsx",
]

def clean(v):
    if v is None:
        return None
    if isinstance(v, str):
        return v.replace("\\n", " ").strip()
    return v

def sheet_snapshot(ws):
    rows = []
    for row in ws.iter_rows():
        vals = [clean(c.value) for c in row]
        if any(v is not None for v in vals):
            rows.append(vals)
    return rows

def inspect(path):
    wb = load_workbook(path, data_only=False, read_only=False)
    out = {"path": str(path.relative_to(ROOT)), "sheets": []}
    for ws in wb.worksheets:
        snap = sheet_snapshot(ws)
        out["sheets"].append({
            "title": ws.title,
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "nonempty_rows": len(snap),
            "rows": snap[:20],
        })
    ws = wb["OGJ2014_Extracted_Rows"]
    headers = [c.value for c in ws[1]]
    out["extracted_headers"] = headers
    source_idx = headers.index("source_report") if "source_report" in headers else None
    if source_idx is not None:
        vals = [ws.cell(r, source_idx+1).value for r in range(2, ws.max_row+1)]
        out["source_report_values"] = vals
        out["source_report_distinct_exact"] = sorted({str(v) for v in vals if v is not None})
    # Search all text cells for provenance/citation vocabulary.
    terms = re.compile(r"doi|journal|author|publication|paper|report|field|reservoir|project|case|table|ogj|oil|gas", re.I)
    hits = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and terms.search(v):
                    hits.append({"sheet": ws.title, "cell": c.coordinate, "value": v})
                    if len(hits) >= 500:
                        break
            if len(hits) >= 500:
                break
        if len(hits) >= 500:
            break
    out["provenance_hits"] = hits
    return out

outs = [inspect(p) for p in paths]
comparison = {
    "same_file_size": paths[0].stat().st_size == paths[1].stat().st_size,
}
# Semantic comparison of all populated cells.
wbs = [load_workbook(p, data_only=False, read_only=False) for p in paths]
comparison["same_sheet_names"] = wbs[0].sheetnames == wbs[1].sheetnames
comparison["sheet_differences"] = []
for name in wbs[0].sheetnames:
    if name not in wbs[1].sheetnames:
        comparison["sheet_differences"].append({"sheet": name, "reason": "missing in copy 2"})
        continue
    a,b=wbs[0][name],wbs[1][name]
    diffs=[]
    mr=max(a.max_row,b.max_row); mc=max(a.max_column,b.max_column)
    for r in range(1,mr+1):
        for c in range(1,mc+1):
            av=a.cell(r,c).value; bv=b.cell(r,c).value
            if av != bv:
                diffs.append({"cell":f"{a.cell(r,c).coordinate}","copy1":av,"copy2":bv})
                if len(diffs)>=100:
                    break
        if len(diffs)>=100: break
    if diffs:
        comparison["sheet_differences"].append({"sheet":name,"differences":diffs})
comparison["identical_semantic_cells"] = not comparison["sheet_differences"]

out = {"comparison": comparison, "copies": outs}
Path(ROOT/"work").mkdir(exist_ok=True)
(Path(ROOT/"work/ogj2014_inspection.json")).write_text(json.dumps(out, indent=2, ensure_ascii=False, default=str), encoding="utf-8")
print(json.dumps(comparison, indent=2))
